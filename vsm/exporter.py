"""Exportación del mapa y sus métricas: PNG, PDF (vectorial) y PowerPoint (.pptx).

* PNG   : imagen del mapa (sin selección ni resaltados de edición) o del gráfico de balance.
* PDF   : informe de 2 páginas A4 apaisado; el mapa se dibuja como vector.
* PPTX  : 4 diapositivas (portada con KPIs, mapa, balance vs takt, tabla de procesos).
          Requiere `pip install python-pptx`.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import List

from PyQt5.QtCore import Qt, QRectF, QPointF, QMarginsF, QByteArray, QBuffer, QIODevice
from PyQt5.QtGui import (QImage, QPainter, QColor, QFont, QPen, QPdfWriter, QPageSize, QPageLayout)

from .balance_chart import balance_data, paint_balance, render_balance_image
from .engine import Metrics

NAVY = QColor("#1f3a6b")
INK = QColor("#2c3e50")
MUTED = QColor("#7f8c8d")
RED = QColor("#c0392b")
HEADER_BG = QColor("#dfe9f5")


class ExportError(Exception):
    """Error legible para mostrar al usuario."""


def default_filename(title: str, ext: str) -> str:
    slug = re.sub(r"[^\w\-]+", "_", title, flags=re.UNICODE).strip("_")[:60] or "mapa_vsm"
    return slug + ext


def _kpis(m: Metrics):
    pce = m.pce * 100
    return [("TAKT TIME", f"{m.takt_time_s:.1f} s" if m.takt_time_s else "—"),
            ("TIEMPO VA", f"{m.va_time_s:g} s"),
            ("LEAD TIME", f"{m.lead_time_days:.2f} días"),
            ("PCE", (f"{pce:.3f} %" if pce < 1 else f"{pce:.1f} %") if m.lead_time_s > 0 else "—"),
            ("CUELLO DE BOTELLA", m.bottleneck_name or "—")]


def _table_rows(m: Metrics) -> List[List[str]]:
    rows = []
    for r in m.process_results:
        pct = f"{r.effective_ct_s / m.takt_time_s * 100:.0f} %" if m.takt_time_s else "—"
        rows.append([r.name, f"{r.cycle_time_s:g}", f"{r.effective_ct_s:.1f}", pct,
                     "Supera el takt" if r.exceeds_takt else "OK"])
    return rows


TABLE_HEADERS = ["Proceso", "TC (s)", "TC efectivo (s)", "% del takt", "Estado"]


def _png_bytes(img: QImage) -> bytes:
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    img.save(buf, "PNG")
    return bytes(ba)


# ================================================================== PNG
def export_png(canvas, path: str, scale: float = 2.0) -> None:
    img = canvas.export_image(scale)
    if img is None:
        raise ExportError("El mapa está vacío.")
    if not img.save(path, "PNG"):
        raise ExportError(f"No se pudo escribir el archivo:\n{path}")


def export_balance_png(m: Metrics, path: str) -> None:
    data = balance_data(m)
    if not data.bars:
        raise ExportError("No hay procesos para graficar el balance.")
    if not render_balance_image(data).save(path, "PNG"):
        raise ExportError(f"No se pudo escribir el archivo:\n{path}")


# ================================================================== PDF
def _font(pt: float, bold: bool = False) -> QFont:
    f = QFont()
    f.setPointSizeF(pt)
    f.setBold(bold)
    return f


def _page_header(p: QPainter, W: float, margin: float, title: str, subtitle: str) -> float:
    p.setPen(NAVY)
    p.setFont(_font(18, True))
    p.drawText(QRectF(margin, margin - 6, W - 2 * margin, 30), Qt.AlignLeft | Qt.AlignVCenter, title)
    p.setPen(MUTED)
    p.setFont(_font(10))
    p.drawText(QRectF(margin, margin + 24, W - 2 * margin, 18), Qt.AlignLeft | Qt.AlignVCenter, subtitle)
    p.setPen(QPen(NAVY, 1.5))
    p.drawLine(QPointF(margin, margin + 48), QPointF(W - margin, margin + 48))
    return margin + 60


def _page_footer(p: QPainter, W: float, H: float, margin: float, n: int) -> None:
    p.setPen(MUTED)
    p.setFont(_font(8))
    p.drawText(QRectF(margin, H - margin + 4, W - 2 * margin, 16), Qt.AlignLeft | Qt.AlignVCenter,
               f"Simulador VSM · {datetime.now():%d/%m/%Y %H:%M}")
    p.drawText(QRectF(margin, H - margin + 4, W - 2 * margin, 16), Qt.AlignRight | Qt.AlignVCenter, f"Página {n}")


def export_pdf(canvas, m: Metrics, title: str, path: str) -> None:
    src = canvas.export_rect()
    if src.isNull():
        raise ExportError("El mapa está vacío.")
    w = QPdfWriter(path)
    w.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    w.setResolution(96)      # 96 dpi = misma proporción texto/símbolo que en pantalla (el PDF sigue siendo vectorial)
    w.setTitle(title)
    w.setCreator("Simulador VSM")
    p = QPainter(w)
    if not p.isActive():
        raise ExportError(f"No se pudo escribir el archivo:\n{path}")
    try:
        p.setRenderHints(QPainter.Antialiasing | QPainter.TextAntialiasing)
        W, H, mg = float(w.width()), float(w.height()), 40.0

        # ---- página 1: mapa
        y = _page_header(p, W, mg, title, "Mapa de flujo de valor · estado actual")
        canvas.render_to(p, QRectF(mg, y, W - 2 * mg, H - mg - y - 6), src)
        _page_footer(p, W, H, mg, 1)

        # ---- página 2: métricas + balance + tabla
        w.newPage()
        y = _page_header(p, W, mg, title, "Métricas y balance de línea (Yamazumi)")
        kp = _kpis(m)
        gap = 10.0
        bw = (W - 2 * mg - gap * (len(kp) - 1)) / len(kp)
        for i, (lab, val) in enumerate(kp):
            box = QRectF(mg + i * (bw + gap), y, bw, 58)
            p.setPen(QPen(QColor("#cfd8dc"), 1))
            p.setBrush(QColor("#ffffff"))
            p.drawRoundedRect(box, 6, 6)
            p.setPen(MUTED)
            p.setFont(_font(7))
            p.drawText(box.adjusted(10, 6, -6, 0), Qt.AlignLeft | Qt.AlignTop, lab)
            p.setPen(RED if (i == 4 and m.processes_over_takt) else INK)
            p.setFont(_font(13, True))
            p.drawText(box.adjusted(10, 22, -6, -4), Qt.AlignLeft | Qt.AlignVCenter, val)
        y += 70
        chart = QRectF(mg, y, W - 2 * mg, 300)
        paint_balance(p, chart, balance_data(m), k=chart.width() / 800.0)
        y = chart.bottom() + 12

        rows = _table_rows(m)
        avail = H - mg - y - 6
        rh = max(14.0, min(22.0, avail / (len(rows) + 1)))
        fracs = [0.34, 0.14, 0.18, 0.16, 0.18]
        widths = [f * (W - 2 * mg) for f in fracs]

        def draw_row(yy, cells, bold=False, bg=None, color=INK):
            x = mg
            if bg is not None:
                p.setPen(Qt.NoPen)
                p.setBrush(bg)
                p.drawRect(QRectF(mg, yy, W - 2 * mg, rh))
            p.setPen(color)
            p.setFont(_font(9, bold))
            for j, (cell, cw) in enumerate(zip(cells, widths)):
                fl = Qt.AlignLeft if j in (0, 4) else Qt.AlignRight
                p.drawText(QRectF(x + 6, yy, cw - 12, rh), fl | Qt.AlignVCenter, cell)
                x += cw

        # la tabla se reparte en varias páginas si hay muchos procesos
        first = max(1, int(avail // rh) - 1)
        cont = max(1, int((H - mg - (mg + 60) - 6) // rh) - 1)
        chunks, rest = [rows[:first]], rows[first:]
        while rest:
            chunks.append(rest[:cont])
            rest = rest[cont:]
        page = 2
        for n, chunk in enumerate(chunks):
            if n:
                _page_footer(p, W, H, mg, page)
                w.newPage()
                page += 1
                y = _page_header(p, W, mg, title, "Tiempos de ciclo por proceso (continuación)")
            draw_row(y, TABLE_HEADERS, True, HEADER_BG, NAVY)
            for i, r in enumerate(chunk):
                over = r[4] != "OK"
                draw_row(y + rh * (i + 1), r, over,
                         QColor("#fdecea") if over else (QColor("#f7f9fb") if i % 2 else None),
                         RED if over else INK)
        _page_footer(p, W, H, mg, page)
    finally:
        p.end()


# ================================================================== PPTX
def export_pptx(canvas, m: Metrics, title: str, path: str) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ExportError("Falta la librería python-pptx.\nInstálala con:  pip install python-pptx") from exc

    map_img = canvas.export_image(2.0)
    if map_img is None:
        raise ExportError("El mapa está vacío.")
    chart_img = render_balance_image(balance_data(m), 1800, 800)

    def rgb(c: QColor) -> RGBColor:
        return RGBColor(c.red(), c.green(), c.blue())

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, s, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        para.alignment = align
        run = para.add_run()
        run.text = s
        run.font.size, run.font.bold = Pt(size), bold
        run.font.color.rgb = rgb(color)
        return tb

    def picture(slide, img, x, y, w, h):
        s = min(w / img.width(), h / img.height())
        pw, ph = img.width() * s, img.height() * s
        slide.shapes.add_picture(io.BytesIO(_png_bytes(img)), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                 Inches(pw), Inches(ph))

    def slide_with_title(t):
        s = prs.slides.add_slide(blank)
        text(s, 0.5, 0.35, 12.3, 0.8, t, 28, True, NAVY)
        return s

    # 1. portada + KPIs
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    text(s, 0.8, 1.6, 11.7, 1.6, title, 40, True, QColor("#ffffff"))
    text(s, 0.8, 3.2, 11.7, 0.6, f"Mapa de flujo de valor · estado actual · {datetime.now():%d/%m/%Y}", 18,
         False, QColor("#bcd0f5"))
    kp = _kpis(m)
    bw, gap = 2.25, 0.2
    for i, (lab, val) in enumerate(kp):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * (bw + gap)), Inches(4.7),
                                 Inches(bw), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(QColor("#2b4d8c"))
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = lab
        for r in tf.paragraphs[0].runs:
            r.font.size, r.font.color.rgb = Pt(11), rgb(QColor("#bcd0f5"))
        para = tf.add_paragraph()
        run = para.add_run()
        run.text = val
        run.font.size, run.font.bold, run.font.color.rgb = Pt(20 if len(val) < 14 else 15), True, RGBColor(255, 255, 255)

    # 2. mapa
    s = slide_with_title("Mapa del estado actual")
    picture(s, map_img, 0.5, 1.25, 12.33, 5.75)

    # 3. balance
    s = slide_with_title("Balance de línea (Yamazumi) vs Takt Time")
    picture(s, chart_img, 0.5, 1.2, 12.33, 5.3)
    data = balance_data(m)
    over = ", ".join(r.name for r in m.processes_over_takt) or "ninguno"
    text(s, 0.5, 6.6, 12.33, 0.6,
         f"Superan el takt ({m.takt_time_s:.1f} s): {over}  ·  Eficiencia de balance: {data.balance_efficiency * 100:.0f} %",
         14, False, MUTED)

    # 4. tabla
    s = slide_with_title("Tiempos de ciclo por proceso")
    rows = _table_rows(m)
    fs = 14 if len(rows) <= 9 else 10
    rh_in = 0.42 if len(rows) <= 9 else max(0.26, min(0.42, 5.8 / (len(rows) + 1)))
    shape = s.shapes.add_table(len(rows) + 1, len(TABLE_HEADERS), Inches(0.5), Inches(1.3), Inches(12.33),
                               Inches(rh_in * (len(rows) + 1)))
    tbl = shape.table
    for j, wd in enumerate((4.6, 1.7, 2.3, 1.8, 1.93)):
        tbl.columns[j].width = Inches(wd)

    def cell(i, j, val, bold=False, color=None):
        c = tbl.cell(i, j)
        c.text_frame.text = val
        para = c.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if j in (0, 4) else PP_ALIGN.RIGHT
        for r in para.runs:
            r.font.size, r.font.bold = Pt(fs), bold
            if color is not None:
                r.font.color.rgb = rgb(color)

    for j, h in enumerate(TABLE_HEADERS):
        cell(0, j, h, True)
    for i, r in enumerate(rows, start=1):
        over = r[4] != "OK"
        for j, v in enumerate(r):
            cell(i, j, v, over, RED if over else None)
    prs.save(path)
